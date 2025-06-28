import openpyxl

def append_experiment_result(file_path, experiment_data):
    try:
        workbook = openpyxl.load_workbook(file_path)
    except FileNotFoundError:
        workbook = openpyxl.Workbook()

    sheet = workbook.active

    if sheet['A1'].value is None:
        sheet['A1'] = 'CLIP'
        sheet['B1'] = 'VIT'
        sheet['C1'] = 'MODEL'
        sheet['D1'] = 'Dataset'
        sheet['E1'] = 'ClsTokenLambda'
        sheet['F1'] = 'FeatureClsTokenLambda'
        sheet['G1'] = 'ProbThd'
        sheet['H1'] = 'Scale'
        sheet['I1'] = 'function'
        sheet['J1'] = 'aAcc'
        sheet['K1'] = 'mIoU'
        sheet['L1'] = 'mAcc'

    last_row = sheet.max_row

    for index, result in enumerate(experiment_data, start=1):
        sheet.cell(row=last_row + index, column=1, value=result['CLIP'])
        sheet.cell(row=last_row + index, column=2, value=result['VIT'])
        sheet.cell(row=last_row + index, column=3, value=result['MODEL'])
        sheet.cell(row=last_row + index, column=4, value=result['Dataset'])
        sheet.cell(row=last_row + index, column=5, value=result['ClsTokenLambda'])
        sheet.cell(row=last_row + index, column=6, value=result['FeatureClsTokenLambda'])
        sheet.cell(row=last_row + index, column=7, value=result['ProbThd'])
        sheet.cell(row=last_row + index, column=8, value=str(result['Scale']))
        sheet.cell(row=last_row + index, column=9, value=result['function'])
        sheet.cell(row=last_row + index, column=10, value=result['aAcc'])
        sheet.cell(row=last_row + index, column=11, value=result['mIoU'])
        sheet.cell(row=last_row + index, column=12, value=result['mAcc'])

    workbook.save(file_path)

